#!/usr/bin/env python3
"""Genera el Excel y la copia diaria de la tabla en CSV y PDF.

    python3 reportes.py

Escribe:
  importacion-cochabamba.xlsx   hojas Hoy (con fórmulas), Historial, Dólar, Movimientos y Supuestos
  historial/AAAA-MM-DD.csv      la tabla de hoy (una copia por día; las de días anteriores no se tocan)
  historial/AAAA-MM-DD.pdf      la misma tabla en PDF
  historial/historial.csv       el historial completo, día por día y producto por producto
"""
import csv
import math
import os
from datetime import date, datetime

from comparar import A_TAX_DEF, B_MIN_KG, B_PASO_KG, B_USD_KG, A_USD_UNIDAD, CAJA_EXTERIOR_G, CONVIENE, JUSTO, leer, locales, oferta

DIR = os.path.dirname(os.path.abspath(__file__))
HIST = os.path.join(DIR, "historial")
UNIDADES = 2
COLUMNAS = ["Producto", "Tipo", "Link del producto", "eBay nuevo USD", "Envío EE.UU. USD", "eBay usado USD",
            "Marketplace Bs", "Tiendas Bs", "Dólar BCB", "Dólar paralelo", "Peso g", "Courier /u USD",
            "Costo puesto Bs", "Mejor local Bs", "Ganancia /u Bs", "Margen", "Estado"]


def tipo(nombre):
    n = nombre.lower()
    return "RAM" if "ddr" in n else "SSD" if "ssd" in n else "CPU"


def num(x):
    return float(x) if x not in ("", None) else None


def dolar_del_dia(fx, fecha):
    """Dólar BCB, paralelo (promedio compra/venta) y usado en cálculos de ese día."""
    f = next((r for r in fx if r["fecha"] == fecha), None)
    if not f:
        return None, None, None
    c, v = num(f["paralelo_compra"]), num(f["paralelo_venta"])
    return num(f["oficial"]), (round((c + v) / 2, 3) if c and v else None), num(f["usado"])


def minimo(filas, mkt):
    vals = [float(f["precio_bs"]) for f in filas if ("Marketplace" in f["fuente"]) == mkt]
    return min(vals) if vals else None


def fila_hoy(v, eeuu, cbba, bcb, par, tc):
    prod, peso = v["producto"], float(v["peso_g_con_caja"])
    nw, us = oferta(eeuu, prod, "nuevo"), oferta(eeuu, prod, "usado")
    loc = locales(cbba, prod)
    mk, sh = minimo(loc, True), minimo(loc, False)
    link = (nw or us or {}).get("link", "")
    r = {"Producto": prod, "Tipo": tipo(prod), "Link del producto": link,
         "eBay nuevo USD": num(nw["precio_usd"]) if nw else None, "Envío EE.UU. USD": num(nw["envio_usd"]) if nw else None,
         "eBay usado USD": num(us["precio_usd"]) + num(us["envio_usd"] or 0) if us else None,
         "Marketplace Bs": mk, "Tiendas Bs": sh, "Dólar BCB": bcb, "Dólar paralelo": par, "Peso g": peso}
    mejor = min([x for x in (mk, sh) if x is not None], default=None)
    r["Mejor local Bs"] = mejor
    if nw:
        p, e = r["eBay nuevo USD"], r["Envío EE.UU. USD"] or 0
        kg = max(B_MIN_KG, math.ceil((peso * UNIDADES + CAJA_EXTERIOR_G) / 1000 / B_PASO_KG - 1e-9) * B_PASO_KG)
        courier = min(kg * B_USD_KG / UNIDADES, A_USD_UNIDAD + p * A_TAX_DEF)
        r["Courier /u USD"] = round(courier, 2)
        r["Costo puesto Bs"] = round((p + e + courier) * tc, 2)
    if nw and mejor:
        g = mejor - r["Costo puesto Bs"]
        r["Ganancia /u Bs"], r["Margen"] = round(g, 2), round(g / mejor, 4)
        r["Estado"] = "Conviene" if g / mejor >= CONVIENE else "Justo" if g / mejor >= JUSTO else "No conviene"
    else:
        r["Estado"] = "Sin datos" if not nw and not mejor else "Sin precio eBay" if not nw else "Sin precio local"
    return r


def historial(vigilar, eeuu, cbba, fx):
    """Una fila por producto y día con lo que se vio ese día, y la diferencia contra el día anterior con datos."""
    filas = []
    for v in vigilar:
        prod = v["producto"]
        fechas = sorted({f["fecha"] for f in eeuu + cbba if f["producto"] == prod})
        prev = None
        for fe in fechas:
            ee = [f for f in eeuu if f["producto"] == prod and f["fecha"] == fe]
            lo = [f for f in locales(cbba, prod) if f["fecha"] == fe]
            ult = {}
            for f in ee:
                ult[f["condicion"]] = f
            tot = lambda f: round(float(f["precio_usd"]) + float(f["envio_usd"] or 0), 2) if f else None
            bcb, par, _ = dolar_del_dia(fx, fe)
            r = {"Fecha": fe, "Producto": prod, "Tipo": tipo(prod), "eBay nuevo USD": tot(ult.get("nuevo")),
                 "eBay usado USD": tot(ult.get("usado")), "Marketplace Bs": minimo(lo, True), "Tiendas Bs": minimo(lo, False),
                 "Dólar BCB": bcb, "Dólar paralelo": par,
                 "Link del producto": (ult.get("nuevo") or ult.get("usado") or (lo[0] if lo else {})).get("link", ""),
                 "Fecha anterior": prev["Fecha"] if prev else ""}
            for k in ("eBay nuevo USD", "eBay usado USD", "Marketplace Bs", "Tiendas Bs"):
                a = prev.get(k) if prev else None
                r["Δ " + k] = round(r[k] - a, 2) if r[k] is not None and a is not None else None
            filas.append(r)
            prev = r
    filas.sort(key=lambda r: (r["Fecha"], r["Producto"]), reverse=True)
    return filas


def escribir_csv(ruta, columnas, filas):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as f:  # utf-8-sig: Excel abre bien los acentos
        w = csv.writer(f, lineterminator="\n")
        w.writerow(columnas)
        for r in filas:
            w.writerow(["" if r.get(c) is None else r.get(c) for c in columnas])


def escribir_pdf(ruta, filas, hoy, tc, bcb, par):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(ruta, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm, title=f"Importación Cochabamba {hoy}")
    f2 = lambda x, d=2: "" if x is None else f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    cols = ["Producto", "Tipo", "eBay nuevo USD", "Envío", "eBay usado USD", "Marketplace Bs", "Tiendas Bs",
            "Dólar BCB", "Dólar paralelo", "Courier /u USD", "Costo puesto Bs", "Mejor local Bs", "Ganancia /u Bs",
            "Margen", "Estado"]
    data = [cols]
    for r in filas:
        data.append([Paragraph(r["Producto"], st["BodyText"].clone("p", fontSize=7, leading=8)), r["Tipo"],
                     f2(r["eBay nuevo USD"]), f2(r["Envío EE.UU. USD"]), f2(r["eBay usado USD"]),
                     f2(r["Marketplace Bs"], 0), f2(r["Tiendas Bs"], 0), f2(r["Dólar BCB"]), f2(r["Dólar paralelo"]),
                     f2(r.get("Courier /u USD")), f2(r.get("Costo puesto Bs"), 0), f2(r["Mejor local Bs"], 0),
                     f2(r.get("Ganancia /u Bs"), 0), "" if r.get("Margen") is None else f"{r['Margen'] * 100:.0f} %",
                     r["Estado"]])
    t = Table(data, repeatRows=1, colWidths=[52 * mm] + [None] * (len(cols) - 1))
    estilo = [("FONT", (0, 0), (-1, -1), "Helvetica", 7), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 6.5),
              ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#15212B")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
              ("ALIGN", (2, 0), (-2, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
              ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9DFE5"))]
    color = {"Conviene": "#DDF1E6", "Justo": "#F6EBCB", "No conviene": "#F6DEDE"}
    for i, r in enumerate(filas, 1):
        if r["Estado"] in color:
            estilo.append(("BACKGROUND", (-1, i), (-1, i), colors.HexColor(color[r["Estado"]])))
    t.setStyle(TableStyle(estilo))
    nota = (f"Dólar BCB {f2(bcb)} · paralelo {f2(par)} · cálculos con {f2(tc)} Bs/USD. Pedido de {UNIDADES} unidades, "
            "courier más barato (B: USD 28/kg, mín. 1 kg, saltos de 0,5 kg + 300 g de caja; A: USD 10/u + 7 % tax). "
            "Aduana: incluida en los dos couriers (no se suma). Estado: conviene ≥ 15 %, justo 5–15 %, no conviene < 5 % sobre el mejor precio local.")
    doc.build([Paragraph(f"Importación EE.UU. → Cochabamba · {hoy}", st["Title"]), Paragraph(nota, st["BodyText"]),
               Spacer(1, 4 * mm), t])


def escribir_excel(ruta, vigilar, eeuu, cbba, fx, hoy_filas, hist):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    F = "Arial"
    azul, negro = Font(name=F, color="0000FF"), Font(name=F)
    head = Font(name=F, bold=True, color="FFFFFF")
    headfill = PatternFill("solid", fgColor="15212B")
    amarillo = PatternFill("solid", fgColor="FFFF00")

    def cabecera(ws, fila, cols, anchos):
        for i, (c, a) in enumerate(zip(cols, anchos), 1):
            cell = ws.cell(fila, i, c)
            cell.font, cell.fill = head, headfill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = a
        ws.row_dimensions[fila].height = 30
        ws.freeze_panes = ws.cell(fila + 1, 2)

    # ---------- Supuestos ----------
    su = wb.active
    su.title = "Supuestos"
    su["A1"], su["A1"].font = "Supuestos del cálculo (celdas azules: se pueden cambiar)", Font(name=F, bold=True, size=13)
    sup = [("Unidades por pedido", UNIDADES, "0"), ("Courier B: USD por kg", B_USD_KG, "0.00"),
           ("Courier B: mínimo kg", B_MIN_KG, "0.0"), ("Courier B: salto kg", B_PASO_KG, "0.0"),
           ("Caja exterior + relleno, g", CAJA_EXTERIOR_G, "0"), ("Courier A: USD por producto", A_USD_UNIDAD, "0.00"),
           ("Courier A: sales tax", A_TAX_DEF, "0.0%"), ("Aduana boliviana extra (incluida en los couriers)", 0, "0.0%"),
           ("Dólar para cálculos, Bs/USD", "=INDEX(Dólar!$F:$F,MATCH(MAX(Dólar!$A:$A),Dólar!$A:$A,0))", "0.00"),
           ("Umbral 'Conviene'", CONVIENE, "0%"), ("Umbral 'Justo'", JUSTO, "0%")]
    for i, (k, v, fmtn) in enumerate(sup, 3):
        su.cell(i, 1, k).font = negro
        c = su.cell(i, 2, v)
        c.font = negro if isinstance(v, str) else azul
        c.number_format = fmtn
    su["B10"].fill = amarillo
    su["B10"].comment = Comment("Los dos couriers ya incluyen la aduana (Santi, 13/10): queda en 0. Probar aquí solo si algún envío la cobra aparte, por ejemplo 0,1 = 10 %.", "Agente")
    su["B11"].comment = Comment("Último 'usado en cálculos' de la hoja Dólar.", "Agente")
    su["A15"] = "Fuente de los supuestos: CLAUDE.md (couriers del usuario). Dólar: hoja Dólar (BCB y paralelo P2P)."
    su["A15"].font = Font(name=F, italic=True, color="5B6873")
    su.column_dimensions["A"].width, su.column_dimensions["B"].width = 36, 14

    # ---------- Dólar ----------
    dl = wb.create_sheet("Dólar")
    cabecera(dl, 1, ["Fecha", "Oficial BCB", "Paralelo compra", "Paralelo venta", "Paralelo promedio", "Usado en cálculos", "Fuente"],
             [12, 13, 15, 15, 16, 16, 80])
    for i, f in enumerate(fx, 2):
        dl.cell(i, 1, datetime.strptime(f["fecha"], "%Y-%m-%d")).number_format = "dd/mm/yyyy"
        for j, k in ((2, "oficial"), (3, "paralelo_compra"), (4, "paralelo_venta"), (6, "usado")):
            if f[k]:
                c = dl.cell(i, j, float(f[k]))
                c.font, c.number_format = azul, "0.00"
        c = dl.cell(i, 5, f"=IF(COUNT(C{i}:D{i})=2,AVERAGE(C{i}:D{i}),\"\")")
        c.font, c.number_format = negro, "0.000"
        dl.cell(i, 7, f["fuente"]).font = negro

    # ---------- Hoy ----------
    ws = wb.create_sheet("Hoy", 0)
    ws["A1"] = f"¿Qué conviene traer? · {date.today().strftime('%d/%m/%Y')}"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = ("eBay: vendedor en EE.UU., ≥ 98 % de feedback, nuevo sin open box. Marketplace y Tiendas: el precio más bajo "
                "que cuenta en Cochabamba. Costo puesto con el courier más barato, según la hoja Supuestos.")
    ws["A2"].font = Font(name=F, italic=True, color="5B6873")
    cabecera(ws, 4, COLUMNAS, [40, 7, 16, 12, 11, 12, 13, 12, 10, 11, 8, 12, 13, 13, 13, 9, 15])
    for i, r in enumerate(hoy_filas, 5):
        ws.cell(i, 1, r["Producto"]).font = negro
        ws.cell(i, 2, r["Tipo"]).font = negro
        if r["Link del producto"]:
            c = ws.cell(i, 3, "ver en eBay")
            c.hyperlink, c.font = r["Link del producto"], Font(name=F, color="C9701C", underline="single")
        for col, k, fmtn in ((4, "eBay nuevo USD", "#,##0.00"), (5, "Envío EE.UU. USD", "#,##0.00"), (6, "eBay usado USD", "#,##0.00"),
                             (7, "Marketplace Bs", "#,##0"), (8, "Tiendas Bs", "#,##0"), (11, "Peso g", "#,##0")):
            if r[k] is not None:
                c = ws.cell(i, col, r[k])
                c.font, c.number_format = azul, fmtn
        ws.cell(i, 9, "=INDEX(Dólar!$B:$B,MATCH(MAX(Dólar!$A:$A),Dólar!$A:$A,0))").number_format = "0.00"
        ws.cell(i, 10, "=INDEX(Dólar!$E:$E,MATCH(MAX(Dólar!$A:$A),Dólar!$A:$A,0))").number_format = "0.000"
        ws.cell(i, 12, f'=IF(D{i}="","",MIN(MAX(Supuestos!$B$5,CEILING((K{i}*Supuestos!$B$3+Supuestos!$B$7)/1000,Supuestos!$B$6))'
                       f'*Supuestos!$B$4/Supuestos!$B$3,Supuestos!$B$8+D{i}*Supuestos!$B$9))').number_format = "#,##0.00"
        ws.cell(i, 13, f'=IF(D{i}="","",(D{i}+E{i}+L{i}+(D{i}+E{i})*Supuestos!$B$10)*Supuestos!$B$11)').number_format = "#,##0"
        ws.cell(i, 14, f'=IF(COUNT(G{i}:H{i})=0,"",MIN(G{i}:H{i}))').number_format = "#,##0"
        ws.cell(i, 15, f'=IF(OR(M{i}="",N{i}=""),"",N{i}-M{i})').number_format = "#,##0;-#,##0;-"
        ws.cell(i, 16, f'=IF(O{i}="","",O{i}/N{i})').number_format = "0%"
        ws.cell(i, 17, f'=IF(P{i}="",IF(AND(D{i}="",N{i}=""),"Sin datos",IF(D{i}="","Sin precio eBay","Sin precio local")),'
                       f'IF(P{i}>=Supuestos!$B$12,"Conviene",IF(P{i}>=Supuestos!$B$13,"Justo","No conviene")))')
        for col in (9, 10, 12, 13, 14, 15, 16, 17):
            ws.cell(i, col).font = negro
    fin = 4 + len(hoy_filas)
    ws.add_table(Table(displayName="Hoy", ref=f"A4:Q{fin}", tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))
    from openpyxl.formatting.rule import CellIsRule
    for txt, color in (("Conviene", "DDF1E6"), ("Justo", "F6EBCB"), ("No conviene", "F6DEDE")):
        ws.conditional_formatting.add(f"Q5:Q{fin}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=PatternFill("solid", fgColor=color)))
    ws.cell(fin + 2, 1, "Azul: datos de eBay, Marketplace y tiendas (de los CSV del repo). Negro: fórmulas. "
                        "Para probar otros supuestos (unidades, aduana, dólar) cambia la hoja Supuestos.").font = Font(name=F, italic=True, color="5B6873")

    # ---------- Historial ----------
    hs = wb.create_sheet("Historial", 1)
    hcols = ["Fecha", "Producto", "Tipo", "eBay nuevo USD", "eBay usado USD", "Marketplace Bs", "Tiendas Bs", "Dólar BCB",
             "Dólar paralelo", "Link del producto", "Fecha anterior", "Δ eBay nuevo", "Δ eBay usado", "Δ Marketplace", "Δ Tiendas"]
    cabecera(hs, 1, hcols, [12, 40, 7, 13, 13, 13, 12, 10, 11, 16, 13, 12, 12, 13, 11])
    for i, r in enumerate(hist, 2):
        hs.cell(i, 1, datetime.strptime(r["Fecha"], "%Y-%m-%d")).number_format = "dd/mm/yyyy"
        hs.cell(i, 2, r["Producto"]); hs.cell(i, 3, r["Tipo"])
        for col, k, fmtn in ((4, "eBay nuevo USD", "#,##0.00"), (5, "eBay usado USD", "#,##0.00"), (6, "Marketplace Bs", "#,##0"), (7, "Tiendas Bs", "#,##0")):
            if r[k] is not None:
                c = hs.cell(i, col, r[k]); c.font, c.number_format = azul, fmtn
        hs.cell(i, 8, f'=IFERROR(IF(INDEX(Dólar!$B:$B,MATCH(A{i},Dólar!$A:$A,0))="","",INDEX(Dólar!$B:$B,MATCH(A{i},Dólar!$A:$A,0))),"")').number_format = "0.00"
        hs.cell(i, 9, f'=IFERROR(INDEX(Dólar!$E:$E,MATCH(A{i},Dólar!$A:$A,0)),"")').number_format = "0.000"
        if r["Link del producto"]:
            c = hs.cell(i, 10, "ver anuncio"); c.hyperlink, c.font = r["Link del producto"], Font(name=F, color="C9701C", underline="single")
        if r["Fecha anterior"]:
            hs.cell(i, 11, datetime.strptime(r["Fecha anterior"], "%Y-%m-%d")).number_format = "dd/mm/yyyy"
        for col, src in ((12, "D"), (13, "E"), (14, "F"), (15, "G")):
            hs.cell(i, col, f'=IF(OR($K{i}="",{src}{i}=""),"",IF(COUNTIFS($B:$B,$B{i},$A:$A,$K{i},{src}:{src},">=0")=0,"",'
                            f'{src}{i}-SUMIFS({src}:{src},$B:$B,$B{i},$A:$A,$K{i})))').number_format = "+#,##0.00;-#,##0.00;0"
        for col in range(1, 16):
            if col not in (4, 5, 6, 7, 10):
                hs.cell(i, col).font = negro
    hs.add_table(Table(displayName="Historial", ref=f"A1:O{1 + len(hist)}", tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))
    from openpyxl.formatting.rule import CellIsRule as R
    rojo, verde = PatternFill("solid", fgColor="F6DEDE"), PatternFill("solid", fgColor="DDF1E6")
    rng_ebay, rng_loc = f"L2:M{1 + len(hist)}", f"N2:O{1 + len(hist)}"
    hs.conditional_formatting.add(rng_ebay, R(operator="greaterThan", formula=["0"], fill=rojo))
    hs.conditional_formatting.add(rng_ebay, R(operator="lessThan", formula=["0"], fill=verde))
    hs.conditional_formatting.add(rng_loc, R(operator="greaterThan", formula=["0"], fill=verde))
    hs.conditional_formatting.add(rng_loc, R(operator="lessThan", formula=["0"], fill=rojo))

    # ---------- Movimientos (cada precio visto) ----------
    mv = wb.create_sheet("Movimientos")
    cabecera(mv, 1, ["Fecha", "Origen", "Producto", "Condición / vendedor", "Precio", "Moneda", "Detalle", "Link"], [12, 12, 40, 36, 11, 8, 60, 14])
    i = 2
    for f in eeuu:
        vals = [datetime.strptime(f["fecha"], "%Y-%m-%d"), "eBay", f["producto"], f["condicion"],
                round(float(f["precio_usd"]) + float(f["envio_usd"] or 0), 2), "USD", f["fuente"]]
        for j, x in enumerate(vals, 1):
            mv.cell(i, j, x)
        if f["link"]:
            c = mv.cell(i, 8, "ver"); c.hyperlink, c.font = f["link"], Font(name=F, color="C9701C", underline="single")
        i += 1
    for f in cbba:
        mkt = "Marketplace" in f["fuente"]
        vals = [datetime.strptime(f["fecha"], "%Y-%m-%d"), "Marketplace" if mkt else "Tienda", f["producto"],
                f["fuente"], float(f["precio_bs"]), "Bs", f["estado"]]
        for j, x in enumerate(vals, 1):
            mv.cell(i, j, x)
        if f["link"]:
            c = mv.cell(i, 8, "ver"); c.hyperlink, c.font = f["link"], Font(name=F, color="C9701C", underline="single")
        i += 1
    for r in range(2, i):
        mv.cell(r, 1).number_format = "dd/mm/yyyy"
        mv.cell(r, 5).number_format = "#,##0.00"
        for c in range(1, 8):
            mv.cell(r, c).font = negro
    mv.add_table(Table(displayName="Movimientos", ref=f"A1:H{i - 1}", tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))

    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if c.font.name != F:
                    c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
    wb.calculation.fullCalcOnLoad = True
    wb.save(ruta)


def main():
    vigilar, eeuu, cbba, fx = leer("vigilar.csv"), leer("precios-eeuu.csv"), leer("precios-cochabamba.csv"), leer("tipo-cambio.csv")
    hoy = date.today().isoformat()
    ultimo = max(fx, key=lambda f: f["fecha"])
    bcb, par, tc = dolar_del_dia(fx, ultimo["fecha"])
    hoy_filas = [fila_hoy(v, eeuu, cbba, bcb, par, tc) for v in vigilar]
    orden = {"Conviene": 0, "Justo": 1, "No conviene": 2}
    hoy_filas.sort(key=lambda r: (orden.get(r["Estado"], 3), -(r.get("Margen") or -9)))
    hist = historial(vigilar, eeuu, cbba, fx)

    os.makedirs(HIST, exist_ok=True)
    escribir_csv(os.path.join(HIST, f"{hoy}.csv"), COLUMNAS, hoy_filas)
    escribir_pdf(os.path.join(HIST, f"{hoy}.pdf"), hoy_filas, hoy, tc, bcb, par)
    hcols = ["Fecha", "Producto", "Tipo", "eBay nuevo USD", "eBay usado USD", "Marketplace Bs", "Tiendas Bs", "Dólar BCB",
             "Dólar paralelo", "Link del producto", "Fecha anterior", "Δ eBay nuevo USD", "Δ eBay usado USD", "Δ Marketplace Bs", "Δ Tiendas Bs"]
    escribir_csv(os.path.join(HIST, "historial.csv"), hcols, hist)
    escribir_excel(os.path.join(DIR, "importacion-cochabamba.xlsx"), vigilar, eeuu, cbba, fx, hoy_filas, hist)
    print(f"importacion-cochabamba.xlsx · historial/{hoy}.csv · historial/{hoy}.pdf · historial/historial.csv ({len(hist)} filas)")


if __name__ == "__main__":
    main()
